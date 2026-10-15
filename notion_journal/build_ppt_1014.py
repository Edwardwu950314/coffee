from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

ESP = C(0x2B, 0x1D, 0x14); AMB = C(0xC7, 0x7D, 0x2E); RED = C(0xB0, 0x3A, 0x2E)
INK = C(0x2B, 0x2B, 0x2B); MUT = C(0x6B, 0x6B, 0x6B); TINT = C(0xF3, 0xF0, 0xEC)
WH = C(255, 255, 255); LIGHT = C(0xE8, 0xDC, 0xCB)
FONT = "Microsoft JhengHei"
p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
blank = p.slide_layouts[6]


def bg(s, color):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = color


def box(s, x, y, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, l in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.space_after = Pt(6)
        r = para.add_run(); r.text = l; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = font
    return t


def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: r.adjustments[0] = 0.06
    r.shadow.inherit = False
    return r


def slide(title, dark=False):
    s = p.slides.add_slide(blank); bg(s, ESP if dark else WH)
    box(s, 0.6, 0.4, 12, 0.8, title, 32, WH if dark else ESP, True)
    return s


def pic(s, f, x, y, w=None, h=None):
    return s.shapes.add_picture('img/' + f, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def code(s, x, y, w, h, text, size=12):
    rect(s, x, y, w, h, TINT)
    t = box(s, x + 0.25, y + 0.2, w - 0.5, h - 0.4, text.split('\n'), size, INK, font="Consolas")
    for para in t.text_frame.paragraphs: para.space_after = Pt(0)


def table(s, x, y, w, rows, colw, size=13, rowh=0.42, hl=None):
    sh = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    tb = sh.table
    for j, cw in enumerate(colw): tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(rowh)
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = str(v); c.margin_left = c.margin_right = Inches(0.1)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            f = c.text_frame.paragraphs[0].runs[0].font; f.size = Pt(size); f.name = FONT
            c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = ESP; f.color.rgb = WH; f.bold = True
            elif hl == i: c.fill.fore_color.rgb = LIGHT; f.color.rgb = INK; f.bold = True
            else: c.fill.fore_color.rgb = TINT if i % 2 == 0 else WH; f.color.rgb = INK
    return tb



E = 'errors/'

# 1 title
s = p.slides.add_slide(blank); bg(s, ESP)
box(s, 0.9, 2.2, 11.5, 1.2, "咖啡豆分選機", 54, WH, True)
box(s, 0.9, 3.5, 11.5, 0.8, "週工作進度報告：日曬 / 蜜處理資料集與模型", 30, LIGHT)
box(s, 0.9, 4.6, 11.5, 0.5, "2026/10/08 – 10/14　｜　ResNet18 好豆 / 壞豆二分類，依處理法各一個模型", 18, AMB)

# 2 headline
s = slide("本週成果：三種處理法皆完成第一版模型")
stats = [("99.3%", "日曬 test 準確率", "漏抓 1、誤判 0"), ("99.3%", "蜜處理 test 準確率", "漏抓 1、誤判 0"),
         ("97.0%", "水洗 test 準確率", "漏抓 5、誤判 1"), ("3,373", "三種處理法圖片總數", "水洗＋日曬＋蜜處理")]
for i, (n, l, sub) in enumerate(stats):
    x = 0.6 + i * 3.08; rect(s, x, 1.7, 2.85, 3.3, TINT)
    box(s, x + 0.2, 2.1, 2.45, 1, n, 44, AMB, True)
    box(s, x + 0.2, 3.4, 2.45, 0.5, l, 16, ESP, True)
    box(s, x + 0.2, 3.95, 2.45, 0.8, sub, 13, MUT)
box(s, 0.6, 5.5, 12.1, 1.2, ["新增：壞豆信心度 p_bad，以及每次訓練後自動列出誤判 / 漏抓的照片。",
                             "注意：日曬、蜜處理只在單日拍攝，99.3% 可能偏樂觀，需補拍獨立驗證集。"], 17, INK)

# 3 dataset
s = slide("資料集與切分")
rows = [["處理法", "good", "bad", "合計", "train", "val", "test"],
        ["水洗 water", 680, 666, "1,346", 942, 202, 202],
        ["日曬 sun", 509, 506, "1,015", 711, 152, 152],
        ["蜜處理 honey", 509, 503, "1,012", 710, 151, 151]]
table(s, 0.6, 1.7, 12.1, rows, [2.8, 1.5, 1.5, 1.7, 1.5, 1.55, 1.55], 17, 0.75)
steps = ["逐顆人工挑選、單顆單獨拍攝，標籤由資料夾名稱 good / bad 取得",
         "prepare_dataset.py：偵測豆子位置，裁成 256×256 方形小圖",
         "依「類別 × 拍攝日期」分層抽樣切成 70 / 15 / 15（seed = 42），存成 splits/<method>.json",
         "good : bad 約 1 : 1，不需處理類別不均；日曬、蜜處理皆為 10/14 單日拍攝"]
for i, t in enumerate(steps):
    box(s, 0.6, 5.0 + i * 0.52, 12.1, 0.5, "• " + t, 15, INK)

# 4 config
s = slide("訓練設定（三種處理法相同）")
rows = [["項目", "設定"], ["模型", "ResNet18（ImageNet 預訓練）＋新的線性分類頭"],
        ["輸入", "256×256 單顆豆子圖 → 224×224"], ["Optimizer", "AdamW，lr 1e-4，weight decay 0.05"],
        ["Scheduler", "ReduceLROnPlateau：val loss 連續 3 epoch 無進步 → lr 減半"],
        ["Loss", "CrossEntropy，label smoothing 0.1，類別權重"], ["正則化", "Dropout 0.5（分類頭前）"],
        ["Batch / Epoch", "batch 16，最多 30 epoch，early stopping patience 8"],
        ["選模型", "以 val macro-F1 挑最佳 epoch（同分取 val loss 較低者）"], ["其他", "seed 42；GPU 自動啟用混合精度 (AMP)"]]
table(s, 0.6, 1.6, 12.1, rows, [2.6, 9.5], 16, 0.55)

# 5 code train
s = slide("程式碼 ①　train.py / dataset.py：訓練核心與資料增強")
code(s, 0.6, 1.6, 6.2, 3.7, '''model = CoffeeBeanClassifier(
    num_classes=2, arch="resnet18", dropout=0.5)
criterion = nn.CrossEntropyLoss(
    weight=class_weights, label_smoothing=0.1)
optimizer = optim.AdamW(
    model.parameters(), lr=1e-4, weight_decay=0.05)
scheduler = optim.lr_scheduler.ReduceLROnPlateau(
    optimizer, mode="min", factor=0.5, patience=3)
# epochs=30, batch_size=16, patience=8, seed=42
# 以 val macro-F1 挑最佳權重''', 13)
code(s, 7.1, 1.6, 5.65, 4.4, '''RandomResizedCrop(224, scale=(0.6, 1.0),
                  ratio=(0.85, 1.15)),
RandomHorizontalFlip(),
RandomVerticalFlip(),
RandomRotation(degrees=30),
ColorJitter(brightness=0.3, contrast=0.3,
            saturation=0.2, hue=0.05),
ToTensor(),
Normalize(IMAGENET_MEAN, IMAGENET_STD),
RandomErasing(p=0.25, scale=(0.02, 0.15),
              value="random"),''', 12.5)
box(s, 0.6, 5.6, 6.2, 1.4, ["三項正則化：Dropout 減少特徵依賴、weight decay 限制權重大小、label smoothing 避免過度自信。"], 14, MUT)
box(s, 7.1, 6.2, 5.65, 0.8, "資料增強只用在訓練集；val / test 只做 Resize + 正規化。", 14, MUT)

# 6 code p_bad
s = slide("程式碼 ②　壞豆信心度與誤判 / 漏抓清單")
code(s, 0.6, 1.6, 6.0, 3.4, '''# predict.py — Predictor.predict()
bad_idx = self.class_names.index("bad")
probs = torch.softmax(self.model(x), dim=1)
scores, indices = probs.max(dim=1)
results += [(self.class_names[i], s, b)
    for i, s, b in zip(indices.tolist(),
        scores.tolist(), probs[:, bad_idx].tolist())]
# 回傳 (判定類別, 判定機率, p_bad)''', 12.5)
code(s, 6.9, 1.6, 5.85, 4.2, '''# train.py — analyze_errors()
for (path, label), p in zip(samples, probs):
    pred = int(p.argmax())
    p_bad = float(p[bad_idx])
    writer.writerow([..., p_bad, int(pred == label)])
    if pred != label:
        (missed if label == bad_idx
         else false_alarm).append((path, p_bad))
missed.sort(key=lambda r: r[1])
false_alarm.sort(key=lambda r: -r[1])''', 12.5)
box(s, 0.6, 5.3, 6.0, 1.7, ["p_bad = 模型認為是壞豆的機率，不論判成哪一類都會給。", "可自訂門檻（如 p_bad > 0.3）減少漏抓，或把 0.3~0.7 的豆子交人工複檢。"], 14, MUT)
box(s, 6.9, 6.0, 5.85, 1.0, "漏抓（壞豆判成好豆）比誤判嚴重，所以分開列出。", 14, MUT)

# 7 sun/honey curves + CM
s = slide("日曬 / 蜜處理：訓練曲線與混淆矩陣")
box(s, 0.5, 1.35, 6.1, 0.4, "日曬　最佳 epoch 19，test 99.3%", 16, ESP, True)
box(s, 6.75, 1.35, 6.1, 0.4, "蜜處理　最佳 epoch 21，test 99.3%", 16, AMB, True)
pic(s, 'sun_1014_curves.png', 0.5, 1.8, w=6.1); pic(s, 'honey_1014_curves.png', 6.75, 1.8, w=6.1)
pic(s, 'sun_1014_cm_test.png', 1.6, 4.2, h=3.0); pic(s, 'honey_1014_cm_test.png', 7.85, 4.2, h=3.0)

# 8 water
s = slide("水洗（重新訓練）：訓練曲線與混淆矩陣")
box(s, 0.5, 1.4, 7.2, 0.4, "最佳 epoch 10，val macro-F1 0.985，test 97.0%", 16, ESP, True)
pic(s, 'water_1014_curves.png', 0.5, 1.9, w=7.3)
pic(s, 'water_1014_cm_test.png', 8.2, 1.5, h=4.4)
box(s, 0.5, 4.8, 7.3, 2.2, ["test 202 張判錯 6 張：漏抓 5、誤判 1", "與 10/08 版本（97.5%）差約 1 張，屬隨機波動範圍", "為了產出錯誤清單，用新版 train.py 重訓"], 15, INK)

# 9 comparison
s = slide("三種處理法比較")
cd = CategoryChartData(); cd.categories = ["水洗", "日曬", "蜜處理"]
cd.add_series("Test Accuracy (%)", (97.0, 99.3, 99.3))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6.2), Inches(5.4), cd); ch = gf.chart
ch.has_legend = False; ch.has_title = True; ch.chart_title.text_frame.text = "Test Accuracy (%)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
pl = ch.plots[0]; pl.has_data_labels = True; pl.data_labels.font.size = Pt(14)
pl.data_labels.number_format = '0.0'; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.gap_width = 60
ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = LIGHT
pt = ser.points[0]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = AMB
va = ch.value_axis; va.minimum_scale = 94; va.maximum_scale = 100; va.has_major_gridlines = False; va.tick_labels.font.size = Pt(12)
ch.category_axis.tick_labels.font.size = Pt(13)
rows = [["處理法", "資料量", "test Acc", "漏抓 / 誤判"], ["水洗", "1,346", "97.0%", "5 / 1"], ["日曬", "1,015", "99.3%", "1 / 0"], ["蜜處理", "1,012", "99.3%", "1 / 0"]]
table(s, 7.0, 1.7, 5.8, rows, [1.5, 1.4, 1.4, 1.5], 14, 0.55)
box(s, 7.0, 4.2, 5.8, 2.8, ["日曬、蜜處理反而比水洗準，但不宜直接解讀為比較好分：",
                            "① 兩批只在單一天、單一條件拍攝，train 與 test 條件幾乎相同",
                            "② test 僅約 150~200 張，錯 1 張差約 0.5~0.7 個百分點，差距在統計上不顯著"], 14, INK)

# 10 errors
s = slide("被判錯的照片（test 集）")
items = [("water_missed_0.089_P_20261007_205845.jpg", "水洗 漏抓", "p_bad 0.089"), ("water_missed_0.119_P_20261007_210823.jpg", "水洗 漏抓", "p_bad 0.119"),
         ("water_missed_0.417_P_20261008_151908.jpg", "水洗 漏抓", "p_bad 0.417"), ("water_missed_0.453_P_20261008_145418.jpg", "水洗 漏抓", "p_bad 0.453"),
         ("water_missed_0.490_P_20261007_205655.jpg", "水洗 漏抓", "p_bad 0.490"), ("water_false_alarm_0.768_P_20261008_154404.jpg", "水洗 誤判", "p_bad 0.768"),
         ("sun_missed_0.267_P_20261014_164140.jpg", "日曬 漏抓", "p_bad 0.267"), ("honey_missed_0.405_P_20261014_135247.jpg", "蜜處理 漏抓", "p_bad 0.405")]
for i, (f, a, b) in enumerate(items):
    x = 0.6 + (i % 4) * 3.08; y = 1.35 + (i // 4) * 2.8
    pic(s, E + f, x, y, w=2.0)
    box(s, x, y + 2.05, 2.8, 0.7, [a, b], 13, RED if "漏抓" in a else ESP, True)
box(s, 0.6, 6.95, 12.1, 0.4, "p_bad 越低的漏抓，代表模型越「理直氣壯地錯」；水洗 0.089、0.119 兩張優先人工複查標籤。", 12, MUT)

# 11 problems
s = slide("遇到的問題與檢討")
rect(s, 0.6, 1.8, 12.1, 3.6, TINT)
box(s, 1.0, 2.2, 11.3, 0.8, "日曬、蜜處理僅單日拍攝", 32, RED, True)
box(s, 1.0, 3.4, 11.3, 1.8, ["無法檢查跨日期的光線、擺放差異，99.3% 可能偏樂觀；", "需補拍獨立驗證集。"], 22, INK)

# 12 next
s = slide("下週計劃", dark=True)
plans = ["人工確認水洗高信心度錯誤的標籤（205845、210823 等）", "補拍日曬、蜜處理獨立驗證集（各約 100 張，不同日期 / 光線）",
         "用驗證集 p_bad 分佈調整判定門檻，畫 ROC / PR 曲線；Grad-CAM 檢查模型注意位置",
         "比較 ResNet18 / ResNet50 / EfficientNet-B0；評估處理法自動判斷與 Jetson Nano 部署"]
for i, t in enumerate(plans):
    y = 1.7 + i * 1.3; rect(s, 0.9, y, 0.7, 0.7, AMB, MSO_SHAPE.OVAL)
    box(s, 0.9, y, 0.7, 0.7, str(i + 1), 20, WH, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 1.9, y, 10.5, 0.9, t, 20, WH, anchor=MSO_ANCHOR.MIDDLE)
p.save('weekly_log_2026-10-14.pptx')
