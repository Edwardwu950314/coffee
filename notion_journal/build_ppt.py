from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

ESP = C(0x2B, 0x1D, 0x14); AMB = C(0xC7, 0x7D, 0x2E); RED = C(0xB0, 0x3A, 0x2E)
INK = C(0x2B, 0x2B, 0x2B); MUT = C(0x6B, 0x6B, 0x6B); TINT = C(0xF3, 0xF0, 0xEC)
WH = C(255, 255, 255); LIGHT = C(0xE8, 0xDC, 0xCB)
FONT = "Microsoft JhengHei"
p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
blank = p.slide_layouts[6]


def bg(s, color):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = color


def box(s, x, y, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, l in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.space_after = Pt(6)
        r = para.add_run(); r.text = l; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = font
    return t


def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: r.adjustments[0] = 0.06
    r.shadow.inherit = False
    return r


def slide(title, dark=False):
    s = p.slides.add_slide(blank); bg(s, ESP if dark else WH)
    box(s, 0.6, 0.4, 12, 0.8, title, 32, WH if dark else ESP, True)
    return s


def pic(s, f, x, y, w=None, h=None):
    return s.shapes.add_picture('img/' + f, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def code(s, x, y, w, h, text, size=12):
    rect(s, x, y, w, h, TINT)
    t = box(s, x + 0.25, y + 0.2, w - 0.5, h - 0.4, text.split('\n'), size, INK, font="Consolas")
    for para in t.text_frame.paragraphs: para.space_after = Pt(0)


def table(s, x, y, w, rows, colw, size=13, rowh=0.42, hl=None):
    sh = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    tb = sh.table
    for j, cw in enumerate(colw): tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(rowh)
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = str(v); c.margin_left = c.margin_right = Inches(0.1)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            f = c.text_frame.paragraphs[0].runs[0].font; f.size = Pt(size); f.name = FONT
            c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = ESP; f.color.rgb = WH; f.bold = True
            elif hl == i: c.fill.fore_color.rgb = LIGHT; f.color.rgb = INK; f.bold = True
            else: c.fill.fore_color.rgb = TINT if i % 2 == 0 else WH; f.color.rgb = INK
    return tb


# 1 title
s = p.slides.add_slide(blank); bg(s, ESP)
box(s, 0.9, 2.2, 11.5, 1.2, "咖啡豆缺陷辨識", 54, WH, True)
box(s, 0.9, 3.5, 11.5, 0.8, "週工作進度報告", 30, LIGHT)
box(s, 0.9, 4.6, 11.5, 0.5, "2026/09/24 – 09/30　｜　ResNet50 好豆 / 壞豆二分類", 18, AMB)

# 2 headline
s = slide("本週成果：轉向好豆 / 壞豆二分類")
stats = [("95.5%", "測試集 Accuracy", "94.7% → 95.5%"), ("0.945", "測試集 macro-F1", "0.936 → 0.945"),
         ("9,015", "訓練＋驗證＋測試", "5573 / 1734 / 1708 張"), ("18 → 2", "分類類別數", "合併 USK-Coffee 資料集")]
for i, (n, l, sub) in enumerate(stats):
    x = 0.6 + i * 3.08; rect(s, x, 1.7, 2.85, 3.3, TINT)
    box(s, x + 0.2, 2.1, 2.45, 1, n, 44, AMB, True)
    box(s, x + 0.2, 3.4, 2.45, 0.5, l, 16, ESP, True)
    box(s, x + 0.2, 3.95, 2.45, 0.8, sub, 13, MUT)
box(s, 0.6, 5.5, 12.1, 1, "主要做法：ResNet50 + Dropout 0.5 + weight decay 0.05 + label smoothing 0.1，並加強資料增強，降低 overfitting。", 17, INK)

# 3 goals
s = slide("本週目標與完成狀態")
rows = [["目標", "狀態"], ["擴充資料（併入 USK-Coffee、自拍豆）", "✅ 完成"], ["離線資料擴增（只擴增 train，避免洩漏）", "✅ 完成"],
        ["改為好豆 / 壞豆二分類", "✅ 完成"], ["比較 ResNet18 / ResNet50 / EfficientNet-B0", "✅ 完成"],
        ["降低 overfitting（Dropout、weight decay、label smoothing）", "✅ 完成"], ["統一「好壞」標準（偏黃、銀皮斑塊）", "⏳ 未決定，列入下週"]]
table(s, 0.6, 1.7, 12.1, rows, [8.6, 3.5], 16, 0.62)

# 4 data
s = slide("資料處理")
steps = [("augment.py", "只擴增 train：翻轉、旋轉、裁切、亮度、模糊、雜訊；val / test 不動"),
         ("merge_classes.py", "合併相近類別（Black / Sour / Insect Damage）"),
         ("resize_class.py", "Normal 類別放大成 500×500，避免解析度成為捷徑"),
         ("make_binary_split.py", "依人工標註 binary_labels.csv 產生 bad / good 切分；USK 沿用原切分")]
for i, (a, b) in enumerate(steps):
    y = 1.6 + i * 1.18; rect(s, 0.6, y, 0.6, 0.6, AMB, MSO_SHAPE.OVAL)
    box(s, 0.6, y, 0.6, 0.6, str(i + 1), 18, WH, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 1.45, y - 0.03, 5.6, 0.4, a, 17, ESP, True, font="Consolas"); box(s, 1.45, y + 0.38, 5.6, 0.8, b, 13, MUT)
table(s, 7.5, 1.7, 5.2, [["切分", "good", "bad", "合計"], ["train", 3626, 1947, 5573], ["val", 1199, 535, 1734], ["test", 1207, 501, 1708]], [1.4, 1.3, 1.3, 1.2], 16, 0.7)
box(s, 7.5, 4.8, 5.2, 1, "壞豆約佔 30%，訓練時用類別加權 loss 處理不平衡。", 14, MUT)

# 5 code model
s = slide("程式碼 ①　model.py：多種 backbone + Dropout")
code(s, 0.6, 1.6, 7.4, 5.1, '''def __init__(self, num_classes=18, pretrained=True,
             arch="resnet18", dropout=0.0):
    weights = "DEFAULT" if pretrained else None
    if arch == "efficientnet_b0":
        net = models.efficientnet_b0(weights=weights)
        self.backbone = nn.Sequential(net.features,
                                      net.avgpool)
        feat_dim = net.classifier[1].in_features
    else:  # resnet18 / resnet50 ...
        net = getattr(models, arch)(weights=weights)
        self.backbone = nn.Sequential(
            *list(net.children())[:-1])
        feat_dim = net.fc.in_features

    self.dropout = nn.Dropout(dropout)
    self.classifier = nn.Linear(feat_dim, num_classes)''', 13)
box(s, 8.4, 1.7, 4.4, 4.5, ["重點", "• arch 參數可切換 ResNet18 / ResNet50 / EfficientNet-B0", "• 特徵維度自動取得，不再寫死 512",
                            "• 分類頭前加 Dropout，訓練時隨機丟特徵抑制 overfitting", "• checkpoint 額外存 arch，推論可自動還原"], 15, INK)

# 6 code train
s = slide("程式碼 ②　train.py / dataset.py：正則化與增強")
code(s, 0.6, 1.6, 6.0, 3.0, '''parser.add_argument("--dropout", default=0.5)
parser.add_argument("--weight-decay", default=0.05)
parser.add_argument("--label-smoothing", default=0.1)

criterion = nn.CrossEntropyLoss(
    weight=class_weights,
    label_smoothing=args.label_smoothing)
optimizer = optim.AdamW(model.parameters(),
    lr=args.lr, weight_decay=args.weight_decay)''', 12.5)
code(s, 6.9, 1.6, 5.85, 3.9, '''RandomResizedCrop(size, scale=(0.6, 1.0),
                  ratio=(0.85, 1.15)),
RandomHorizontalFlip(),
RandomVerticalFlip(),
RandomRotation(degrees=30),
ColorJitter(brightness=0.3, contrast=0.3,
            saturation=0.2, hue=0.05),
ToTensor(),
Normalize(IMAGENET_MEAN, IMAGENET_STD),
RandomErasing(p=0.25, scale=(0.02, 0.15),
              value="random"),''', 12.5)
box(s, 0.6, 4.9, 6.0, 1.8, ["三項正則化各司其職：", "Dropout 減少特徵依賴、weight decay 限制權重大小、label smoothing 避免過度自信。"], 14, MUT)
box(s, 6.9, 5.75, 5.85, 1, "增強更強：裁切範圍加大，並新增 RandomErasing。", 14, MUT)

# 7 experiments
s = slide("實驗比較：測試集表現")
cd = CategoryChartData(); cd.categories = ["ResNet18", "EffNet-B0", "ResNet50", "ResNet50 v2", "bin_reg"]
cd.add_series("Test Accuracy (%)", (93.2, 93.2, 94.8, 94.7, 95.5))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6.6), Inches(5.4), cd); ch = gf.chart
ch.has_legend = False; ch.has_title = True; ch.chart_title.text_frame.text = "Test Accuracy (%)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
pl = ch.plots[0]; pl.has_data_labels = True; pl.data_labels.font.size = Pt(13)
pl.data_labels.number_format = '0.0'; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.gap_width = 60
ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = LIGHT
pt = ser.points[4]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = AMB
va = ch.value_axis; va.minimum_scale = 90; va.maximum_scale = 97; va.has_major_gridlines = False; va.tick_labels.font.size = Pt(11)
ch.category_axis.tick_labels.font.size = Pt(11)
rows = [["實驗", "Test Acc", "Test F1"], ["18 類 ResNet18（起點）", "0.611", "0.570"], ["好/壞 ResNet18", "0.932", "0.916"],
        ["好/壞 EffNet-B0", "0.932", "0.917"], ["好/壞 ResNet50", "0.948", "0.935"], ["ResNet50 v2", "0.947", "0.936"], ["bin_reg（正則化）", "0.955", "0.945"]]
table(s, 7.4, 1.6, 5.4, rows, [3.0, 1.2, 1.2], 13, 0.5, hl=6)
box(s, 7.4, 5.3, 5.4, 1.6, ["ResNet50 明顯優於 ResNet18 / EffNet-B0；加入正則化後再進步約 0.8 個百分點。", "18 類起點的任務較難，不能直接與二分類比較。"], 13, MUT)

# 8 curves
s = slide("訓練圖：加入正則化後更平穩")
pic(s, 'v2_curves.png', 0.5, 1.9, w=6.1); pic(s, 'binreg_curves.png', 6.75, 1.9, w=6.1)
box(s, 0.5, 1.45, 6.1, 0.4, "ResNet50 v2（無強正則化）", 16, ESP, True); box(s, 6.75, 1.45, 6.1, 0.4, "bin_reg（Dropout + WD + LS）", 16, AMB, True)
box(s, 0.5, 4.6, 6.1, 1.5, ["Train loss 降到約 0.005，Val loss 在 0.09–0.17 間震盪", "→ 有 overfitting 跡象"], 15, INK)
box(s, 6.75, 4.6, 6.1, 1.5, ["Train loss 停在約 0.24（label smoothing 下限）", "Val loss 更平穩，訓練與驗證差距縮小"], 15, INK)
box(s, 0.5, 6.4, 12.3, 0.6, "起點 18 類 ResNet18（9/23）：Val acc 0.813 / Test acc 0.611，泛化明顯不足。", 14, MUT)

# 9 confusion
s = slide("測試圖：混淆矩陣（test set）")
pic(s, 'v2_cm_test.png', 0.9, 1.65, h=4.4); pic(s, 'binreg_cm_test.png', 7.1, 1.65, h=4.4)
box(s, 0.9, 6.2, 5.3, 0.9, ["ResNet50 v2", "bad→good 誤判 49、good→bad 誤判 42"], 14, INK)
box(s, 7.1, 6.2, 5.3, 0.9, ["bin_reg", "bad→good 誤判 50、good→bad 誤判 27"], 14, INK)

# 10 predictions
s = slide("測試圖：單張預測與分類報告")
pic(s, 'pred_husk.jpg', 0.6, 1.5, h=3.5); pic(s, 'pred_usk.jpg', 3.85, 1.5, h=3.5); pic(s, 'pred_binary.jpg', 7.1, 1.5, h=3.5)
box(s, 0.6, 5.25, 3.1, 0.6, "Husk（瑕疵）→ bad 100%", 12, MUT); box(s, 3.9, 5.25, 3.1, 0.6, "USK 樣本", 12, MUT); box(s, 7.2, 5.25, 3.1, 0.6, "二元模型輸出", 12, MUT)
table(s, 10.85, 1.6, 2.0, [["class", "recall"], ["bad", "0.900"], ["good", "0.978"]], [1.0, 1.0], 13, 0.5)
box(s, 10.85, 3.3, 2.0, 1.8, ["bin_reg 測試集", "bad F1 0.921", "good F1 0.968"], 13, INK)
box(s, 0.6, 6.1, 12.1, 0.9, "batch_predict.py 新增 --per-class 每類抽樣，並支援二元模型（Normal 視為 good）。", 14, MUT)

# 11 problems
s = slide("遇到的問題")
items = [("解析度不一致", "Normal 為 200×200、其他 500×500，可能成為捷徑；已用 resize_class.py 處理"),
         ("標籤矛盾", "Fade_08 與 Partial Sour_06 是同一張圖，切分時已排除"),
         ("好壞標準不一致", "偏黃算不算壞未決定（USK 約 60 張漏抓）；銀皮斑塊只套用在自拍資料"),
         ("bad recall 僅 0.90", "漏抓壞豆（bad→good 50 張）是實務上較嚴重的錯誤")]
for i, (a, b) in enumerate(items):
    x = 0.6 + (i % 2) * 6.15; y = 1.6 + (i // 2) * 2.6; rect(s, x, y, 5.95, 2.3, TINT)
    box(s, x + 0.3, y + 0.3, 5.35, 0.5, a, 20, RED if i == 3 else ESP, True); box(s, x + 0.3, y + 1.0, 5.35, 1.2, b, 15, INK)

# 12 next
s = slide("下週計劃", dark=True)
plans = ["補拍自己的好豆與壞豆（自拍好豆僅 35 張），裁成單顆 500×500 併入 data/merged/", "統一好壞標準，並依新標準複查 USK 好豆",
         "用新資料重訓 ResNet50，分開評估自拍與 USK 的好 / 壞豆表現", "調整決策門檻，提高 bad recall"]
for i, t in enumerate(plans):
    y = 1.7 + i * 1.3; rect(s, 0.9, y, 0.7, 0.7, AMB, MSO_SHAPE.OVAL)
    box(s, 0.9, y, 0.7, 0.7, str(i + 1), 20, WH, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 1.9, y, 10.5, 0.9, t, 20, WH, anchor=MSO_ANCHOR.MIDDLE)
p.save('weekly_log_2026-09-30.pptx')
